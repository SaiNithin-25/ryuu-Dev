import glob
import json
import random
import re
import zipfile
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional

try:
    import pyarrow.parquet as pq
except Exception:
    pq = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document
except Exception:
    Document = None


def load_config(path="data_pipeline/pipeline_config.json"):
    with open(path, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    return cfg


def ensure_parent(path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)


def read_jsonl(path):
    if not Path(path).exists():
        return
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                yield json.loads(line)
            except json.JSONDecodeError:
                continue


def write_jsonl(path, rows):
    ensure_parent(path)
    with open(path, "w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def append_jsonl(path, rows):
    ensure_parent(path)
    with open(path, "a", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def iter_source_files(source_cfg):
    pattern = source_cfg["path"]
    return sorted(glob.glob(pattern, recursive=True))


def iter_rows_from_file(path, fmt="auto"):
    ext = Path(path).suffix.lower()
    file_fmt = fmt
    if fmt == "auto":
        if ext in (".jsonl", ".json"):
            file_fmt = "jsonl"
        elif ext == ".parquet":
            file_fmt = "parquet"
        elif ext == ".pdf":
            file_fmt = "pdf"
        elif ext == ".docx":
            file_fmt = "docx"
        elif ext == ".doc":
            file_fmt = "doc"
        elif ext in (".txt", ".md"):
            file_fmt = "text"
        else:
            return

    if file_fmt == "jsonl":
        yield from read_jsonl(path)
        return

    if file_fmt == "parquet":
        if pq is None:
            return
        pf = pq.ParquetFile(path)
        for batch in pf.iter_batches(batch_size=50000):
            for row in batch.to_pylist():
                yield row
        return

    if file_fmt == "text":
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        yield {"prompt": f"Summarize file {Path(path).name}", "response": text}
        return

    if file_fmt in ("pdf", "docx", "doc"):
        text = extract_document_text(path, file_fmt)
        if not text or not text.strip():
            return
        for i, chunk in enumerate(split_text_chunks(text, chunk_chars=6000, overlap_chars=300), start=1):
            yield {
                "prompt": f"Summarize the technical content from document {Path(path).name}, section {i}.",
                "response": chunk,
            }


PROMPT_KEYS = ["prompt", "instruction", "question", "input", "query", "task"]
RESPONSE_KEYS = ["response", "completion", "output", "answer", "solution", "code"]


def _normalize_text(v: str) -> str:
    v = v.replace("\r\n", "\n").replace("\r", "\n")
    v = re.sub(r"[ \t]+", " ", v)
    v = re.sub(r"\n{3,}", "\n\n", v)
    return v.strip()


def _pick_first_str(row: Dict, keys: List[str]) -> str:
    for k in keys:
        v = row.get(k)
        if isinstance(v, str) and v.strip():
            return _normalize_text(v)
    return ""


def _extract_chat_pair(row: Dict) -> Optional[Dict[str, str]]:
    conv = row.get("conversations") or row.get("messages")
    if not isinstance(conv, list):
        return None
    user_msg = ""
    assistant_msg = ""
    for item in conv:
        if not isinstance(item, dict):
            continue
        role = str(item.get("role", item.get("from", ""))).lower()
        content = item.get("content", item.get("value", ""))
        if not isinstance(content, str) or not content.strip():
            continue
        content = _normalize_text(content)
        if not user_msg and role in ("human", "user"):
            user_msg = content
        elif user_msg and role in ("assistant", "gpt", "bot"):
            assistant_msg = content
            break
    if user_msg and assistant_msg:
        return {"prompt": user_msg, "response": assistant_msg}
    return None


def _join_instruction_input(row: Dict) -> str:
    instruction = row.get("instruction")
    input_text = row.get("input")
    if isinstance(instruction, str) and instruction.strip():
        instruction = _normalize_text(instruction)
        if isinstance(input_text, str) and input_text.strip():
            return f"{instruction}\n\nContext:\n{_normalize_text(input_text)}"
        return instruction
    return ""


def normalize_row(row):
    # Dataset-specific normalization for APPS/OpenMathInstruct
    ds_name = str(row.get("_source", "")).lower()
    if ds_name == "custom":
        # Special case: math_ai_dataset_2M.jsonl schema
        if "problem" in row and ("reasoning" in row or "final_answer" in row):
            return normalize_math_ai_row(row)
    if ds_name == "hf_apps":
        return normalize_apps_row(row)
    if ds_name == "hf_openmath_correct":
        return normalize_openmath_row(row, allow_incorrect=False)
    if ds_name == "hf_openmath_incorrect":
        return normalize_openmath_row(row, allow_incorrect=False)

    chat_pair = _extract_chat_pair(row)
    if chat_pair is not None:
        return chat_pair

    prompt = _pick_first_str(row, PROMPT_KEYS)
    response = _pick_first_str(row, RESPONSE_KEYS)

    if not prompt:
        prompt = _join_instruction_input(row)

    if not response:
        for key in ("solutions", "answers"):
            sols = row.get(key)
            if isinstance(sols, list) and sols:
                first = sols[0]
                if isinstance(first, str) and first.strip():
                    response = _normalize_text(first)
                    break

    if not prompt or not response:
        return None

    return {"prompt": prompt, "response": response}


def normalize_apps_row(row):
    prompt = row.get("question") or row.get("prompt")
    solutions = row.get("solutions")
    response = None

    if isinstance(solutions, str) and solutions.strip():
        try:
            parsed = json.loads(solutions)
            if isinstance(parsed, list) and parsed and isinstance(parsed[0], str):
                response = parsed[0]
        except Exception:
            response = solutions
    elif isinstance(solutions, list) and solutions and isinstance(solutions[0], str):
        response = solutions[0]

    if isinstance(prompt, str) and isinstance(response, str) and prompt.strip() and response.strip():
        return {"prompt": _normalize_text(prompt), "response": _normalize_text(response)}
    return None


def normalize_openmath_row(row, allow_incorrect: bool = False):
    # OpenMathInstruct schema: question + generated_solution (+ expected_answer)
    if not allow_incorrect:
        is_correct = row.get("is_correct")
        if isinstance(is_correct, bool) and not is_correct:
            return None

    prompt = row.get("question") or row.get("problem") or row.get("prompt") or row.get("input")
    response = row.get("generated_solution") or row.get("solution") or row.get("answer") or row.get("final")

    if isinstance(prompt, str) and isinstance(response, str) and prompt.strip() and response.strip():
        return {"prompt": _normalize_text(prompt), "response": _normalize_text(response)}
    return None


def normalize_math_ai_row(row):
    problem = row.get("problem")
    reasoning = row.get("reasoning")
    final_answer = row.get("final_answer")

    if not isinstance(problem, str) or not problem.strip():
        return None

    parts = []
    if isinstance(reasoning, str) and reasoning.strip():
        parts.append(reasoning.strip())
    if final_answer is not None and str(final_answer).strip():
        parts.append(f"Final: {final_answer}")

    if not parts:
        return None

    response = "\n".join(parts)
    return {"prompt": _normalize_text(problem), "response": _normalize_text(response)}


def basic_token_count(text):
    return len(text.split())


def ascii_ratio(text):
    if not text:
        return 1.0
    printable = sum(1 for c in text if ord(c) < 128)
    return printable / max(1, len(text))


def repeat_line_ratio(text):
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return 0.0
    unique = len(set(lines))
    return 1.0 - (unique / max(1, len(lines)))


def url_ratio(text):
    if not text:
        return 0.0
    urls = re.findall(r"https?://\S+|www\.\S+", text)
    if not urls:
        return 0.0
    url_chars = sum(len(u) for u in urls)
    return url_chars / max(1, len(text))


def symbol_ratio(text):
    if not text:
        return 0.0
    sym = sum(1 for c in text if not c.isalnum() and not c.isspace())
    return sym / max(1, len(text))


def unique_words_ratio(text):
    words = re.findall(r"[A-Za-z0-9_]+", text.lower())
    if not words:
        return 0.0
    return len(set(words)) / max(1, len(words))


def canonicalize_text(text):
    text = text.lower().strip()
    text = re.sub(r"\s+", " ", text)
    return text


def set_seed(seed):
    random.seed(seed)


def load_rows(path: str) -> List[Dict]:
    return list(read_jsonl(path) or [])


def split_text_chunks(text: str, chunk_chars: int = 6000, overlap_chars: int = 300) -> Iterator[str]:
    text = _normalize_text(text)
    if not text:
        return
    n = len(text)
    if n <= chunk_chars:
        yield text
        return
    start = 0
    while start < n:
        end = min(n, start + chunk_chars)
        chunk = text[start:end].strip()
        if chunk:
            yield chunk
        if end >= n:
            break
        start = max(start + 1, end - overlap_chars)


def extract_document_text(path: str, file_fmt: str) -> str:
    p = Path(path)
    if file_fmt == "pdf":
        return _extract_pdf_text(p)
    if file_fmt == "docx":
        return _extract_docx_text(p)
    if file_fmt == "doc":
        return _extract_doc_text(p)
    return ""


def _extract_pdf_text(path: Path) -> str:
    if PdfReader is None:
        print(f"[WARN] Skipping {path}: install pypdf to read .pdf files.")
        return ""
    try:
        reader = PdfReader(str(path))
        chunks = []
        for pg in reader.pages:
            txt = pg.extract_text() or ""
            if txt.strip():
                chunks.append(txt)
        return "\n\n".join(chunks)
    except Exception as e:
        print(f"[WARN] PDF read failed for {path}: {e}")
        return ""


def _extract_docx_text(path: Path) -> str:
    if Document is not None:
        try:
            doc = Document(str(path))
            lines = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            return "\n".join(lines)
        except Exception:
            pass

    # Fallback parser for docx (zip+xml) without python-docx.
    try:
        with zipfile.ZipFile(path, "r") as zf:
            data = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        # Replace paragraph boundaries before stripping XML tags.
        data = data.replace("</w:p>", "\n")
        data = re.sub(r"<[^>]+>", "", data)
        return data
    except Exception as e:
        print(f"[WARN] DOCX read failed for {path}: {e}")
        return ""


def _extract_doc_text(path: Path) -> str:
    # Legacy .doc is binary; no robust stdlib parser.
    print(f"[WARN] Skipping {path}: .doc extraction unsupported by default. Convert to .docx or .txt.")
    return ""
